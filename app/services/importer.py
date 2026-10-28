import csv, io, re
from openpyxl import load_workbook

def norm(s):
    s = (s or "").strip().lower()
    return s.replace("é","e").replace("è","e").replace("ê","e").replace("à","a").replace("û","u").replace("ô","o").replace("î","i").replace("ï","i").replace("ë","e").replace(" ", "_")

def parse_guest_file(filename: str, content: bytes):
    if filename.lower().endswith('.csv'):
        text = content.decode('utf-8-sig')
        sample = text[:2048]
        try: dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif filename.lower().endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    else: raise ValueError('Format non pris en charge. Utilisez .xlsx ou .csv.')
    if len(rows) < 2: raise ValueError('Le fichier doit contenir un en-tête et au moins une invitation.')
    headers = [norm(str(x or '')) for x in rows[0]]
    aliases = {
      'name': {'nom','invite','invitee','nom_de_linvite','nom_invite','nom_du_groupe','guest','guest_name'},
      'type': {'type','type_invitation','type_dinvitation'},
      'quota': {'nombre','quota','personnes','nombre_de_personnes','nb_personnes','nombre_autorise'},
      'table': {'table','numero_table','table_attribuee'}
    }
    indexes = {}
    for key, names in aliases.items():
        for i,h in enumerate(headers):
            if h in names: indexes[key] = i; break
    if 'name' not in indexes: raise ValueError('Colonne nom manquante. Utilisez « nom » ou « invite ».')
    if 'quota' not in indexes: raise ValueError('Colonne quota manquante. Utilisez « nombre », « quota » ou « personnes ».')
    guests=[]; errors=[]
    for line_no, row in enumerate(rows[1:], start=2):
        if not row or not any(x not in (None, '') for x in row): continue
        def val(key):
            i=indexes.get(key)
            return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ''
        name=val('name'); rawq=val('quota'); typ=val('type').lower(); table=val('table')
        if not name: errors.append(f'Ligne {line_no}: nom manquant'); continue
        try:
            q=int(float(rawq))
            if q < 1: raise ValueError()
        except Exception: errors.append(f'Ligne {line_no}: quota invalide pour {name}'); continue
        if not typ:
            typ='individuelle' if q==1 else ('couple' if q==2 else 'groupe')
        typ=typ if typ in {'individuelle','individuel','couple','groupe','famille'} else 'groupe'
        if typ=='individuel': typ='individuelle'
        if typ=='famille': typ='groupe'
        guests.append({'guest_name':name,'invitation_type':typ,'quota':q,'table_name':table,'source_line':line_no})
    if errors: raise ValueError('Import annulé; corrigez le fichier : ' + ' ; '.join(errors[:15]))
    if not guests: raise ValueError('Aucune invitation valide trouvée.')
    seen=set(); dup=[]
    for g in guests:
        # FIX: neutraliser les espaces multiples et les espaces de bord pour comparer les noms.
        key=re.sub(r'\s+', ' ', g['guest_name']).strip().casefold()
        if key in seen: dup.append(g['guest_name'])
        seen.add(key)
    return guests, dup
